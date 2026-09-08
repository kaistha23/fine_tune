"""Document extraction: the missing front end of the ingestion pipeline.

`ingest.py` chunked text but nothing produced text from a PDF or Word file, which is how
regulatory circulars actually arrive.

The logic that matters - repeated-header detection, table rendering, page attribution -
is pure Python and runs in the offline container. Only the PDF and Word parsing needs the
`documents` extra, so those tests skip when it is absent.
"""
import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from credit_risk.rag.extract import (
    ExtractionError,
    Page,
    _normalise,
    _render_table,
    extract,
    find_repeated_lines,
    strip_repeated_lines,
)
from credit_risk.rag.ingest import DocumentMeta, chunk_pages
from credit_risk.schemas import Jurisdiction

try:
    import pymupdf
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def meta() -> DocumentMeta:
    return DocumentMeta(
        document_id="SAMA-CIRC-4", jurisdiction=Jurisdiction.SAMA,
        document_version="2.0", approval_status="approved",
        effective_from=date(2025, 1, 1))


HEADER = "SAMA - Confidential - Page {n} of 3"


def sample_pages() -> list[Page]:
    return [
        Page(1, HEADER.format(n=1) + "\n\n1. Scope\n\n"
                "This circular applies to all licensed banks in the Kingdom."),
        Page(2, HEADER.format(n=2) + "\n\n7.2 Significant increase in credit risk\n\n"
                "An exposure moves to stage 2 when creditworthiness has deteriorated "
                "materially since initial recognition."),
        Page(3, HEADER.format(n=3) + "\n\n7.3 Credit impaired\n\n"
                "An exposure more than 90 days past due is classified in stage 3."),
    ]


class RepeatedLineTests(unittest.TestCase):
    def test_a_running_header_is_detected_despite_the_page_number(self) -> None:
        # "Page 1 of 3" and "Page 2 of 3" are the same header. Without normalising the
        # digits every page's header looks unique and none is ever removed.
        repeated = find_repeated_lines(sample_pages())
        self.assertIn("SAMA - Confidential - Page # of #", repeated)

    def test_body_text_is_not_treated_as_a_header(self) -> None:
        repeated = find_repeated_lines(sample_pages())
        self.assertNotIn(_normalise("1. Scope"), repeated)

    def test_stripping_removes_the_header_and_keeps_the_clause(self) -> None:
        pages = strip_repeated_lines(sample_pages(), find_repeated_lines(sample_pages()))
        joined = "\n".join(page.text for page in pages)
        self.assertNotIn("Confidential", joined)
        self.assertIn("7.2 Significant increase in credit risk", joined)

    def test_a_short_document_keeps_everything(self) -> None:
        # Two pages sharing a line is not evidence of a running header, and dropping a
        # clause because it happened to repeat would lose real content.
        pages = [Page(1, "Shared line\n\nBody one"), Page(2, "Shared line\n\nBody two")]
        self.assertEqual(find_repeated_lines(pages), set())

    def test_a_long_line_is_never_a_header(self) -> None:
        long_line = "A " * 100
        pages = [Page(n, long_line) for n in range(1, 6)]
        self.assertEqual(find_repeated_lines(pages), set())


class TableRenderingTests(unittest.TestCase):
    def test_each_row_keeps_its_column_headings(self) -> None:
        # Read as flowing text a provisioning matrix becomes "2 3 3% 100%", which is not
        # recoverable. Row-wise, each number stays attached to what it measures.
        rendered = _render_table([["Stage", "Coverage"], ["2", "3%"], ["3", "100%"]])
        self.assertIn("Stage: 2 | Coverage: 3%", rendered)
        self.assertIn("Stage: 3 | Coverage: 100%", rendered)

    def test_empty_rows_are_dropped(self) -> None:
        rendered = _render_table([["Stage", "Coverage"], ["", ""], ["3", "100%"]])
        self.assertEqual(rendered, "Stage: 3 | Coverage: 100%")

    def test_none_cells_do_not_crash(self) -> None:
        self.assertIn("Stage: 2", _render_table([["Stage", "Coverage"], ["2", None]]))

    def test_an_empty_table_renders_to_nothing(self) -> None:
        self.assertEqual(_render_table([]), "")


class PageAttributionTests(unittest.TestCase):
    def setUp(self) -> None:
        pages = sample_pages()
        self.pages = strip_repeated_lines(pages, find_repeated_lines(pages))
        self.chunks = {c.section_id: c for c in chunk_pages(self.pages, meta())}

    def test_each_clause_carries_the_page_it_starts_on(self) -> None:
        # PolicyChunk.page_number existed and nothing set it, so a citation could name a
        # clause but not where to find it in a ninety-page PDF.
        self.assertEqual(self.chunks["1"].page_number, 1)
        self.assertEqual(self.chunks["7.2"].page_number, 2)
        self.assertEqual(self.chunks["7.3"].page_number, 3)

    def test_chunks_follow_clauses_not_page_breaks(self) -> None:
        # A clause spanning a page break is one rule; splitting it there would strip the
        # condition off it, which is the failure fixed-length chunking causes.
        pages = [
            Page(1, "7.2 Significant increase\n\nAn exposure moves to stage 2 when"),
            Page(2, "creditworthiness has deteriorated materially since recognition."),
        ]
        chunks = chunk_pages(pages, meta())
        text = " ".join(c.text for c in chunks)
        self.assertIn("stage 2 when", text)
        self.assertIn("deteriorated materially", text)
        self.assertEqual(len([c for c in chunks if c.section_id == "7.2"]), 1)

    def test_lineage_survives_extraction(self) -> None:
        for chunk in self.chunks.values():
            self.assertEqual(chunk.document_id, "SAMA-CIRC-4")
            self.assertEqual(chunk.jurisdiction, Jurisdiction.SAMA)
            self.assertTrue(chunk.content_hash)


class UnsupportedFormatTests(unittest.TestCase):
    def test_an_unknown_suffix_is_refused(self) -> None:
        with self.assertRaises(ExtractionError):
            extract("circular.xlsx")

    def test_plain_text_is_read_directly(self) -> None:
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "circular.txt"
            path.write_text("7.2 Staging\n\nStage 2 on SICR.", encoding="utf-8")
            pages = extract(path)
            self.assertEqual(len(pages), 1)
            self.assertIn("Stage 2 on SICR", pages[0].text)


@unittest.skipUnless(HAS_PYMUPDF, "install the 'documents' extra for PDF extraction")
class PdfExtractionTests(unittest.TestCase):
    def _write_pdf(self, directory: Path) -> Path:
        document = pymupdf.open()
        for page in sample_pages():
            new = document.new_page()
            new.insert_text((60, 70), page.text, fontsize=10)
        path = directory / "circular.pdf"
        document.save(str(path))
        document.close()
        return path

    def test_a_pdf_extracts_one_page_at_a_time(self) -> None:
        with TemporaryDirectory() as tmp:
            pages = extract(self._write_pdf(Path(tmp)))
        self.assertEqual([p.page_number for p in pages], [1, 2, 3])
        self.assertIn("Significant increase", pages[1].text)

    def test_the_running_header_is_gone(self) -> None:
        with TemporaryDirectory() as tmp:
            pages = extract(self._write_pdf(Path(tmp)))
        self.assertNotIn("Confidential", "\n".join(p.text for p in pages))

    def test_clauses_reach_the_right_pages_end_to_end(self) -> None:
        with TemporaryDirectory() as tmp:
            chunks = {c.section_id: c
                      for c in chunk_pages(extract(self._write_pdf(Path(tmp))), meta())}
        self.assertEqual(chunks["7.2"].page_number, 2)
        self.assertEqual(chunks["7.3"].page_number, 3)

    def test_a_pdf_with_no_text_is_refused(self) -> None:
        # A scanned circular extracts to nothing. Ingesting it silently would shrink the
        # corpus without anyone noticing the document is missing.
        with TemporaryDirectory() as tmp:
            document = pymupdf.open()
            document.new_page()
            path = Path(tmp) / "scanned.pdf"
            document.save(str(path))
            document.close()
            with self.assertRaises(ExtractionError) as caught:
                extract(path)
        self.assertIn("OCR", str(caught.exception))


@unittest.skipUnless(HAS_DOCX, "install the 'documents' extra for Word extraction")
class DocxExtractionTests(unittest.TestCase):
    def _write_docx(self, directory: Path) -> Path:
        document = docx.Document()
        document.add_heading("Staging", level=1)
        document.add_paragraph("Exposures move to stage 2 on a significant increase.")
        document.add_heading("Provisioning", level=2)
        document.add_paragraph("Stage 3 carries a lifetime expected credit loss.")
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Stage", "Coverage"
        table.cell(1, 0).text, table.cell(1, 1).text = "2", "3%"
        table.cell(2, 0).text, table.cell(2, 1).text = "3", "100%"
        path = directory / "circular.docx"
        document.save(str(path))
        return path

    def test_word_headings_become_section_structure(self) -> None:
        with TemporaryDirectory() as tmp:
            chunks = chunk_pages(extract(self._write_docx(Path(tmp))), meta())
        paths = [c.heading_path for c in chunks]
        self.assertIn(["Staging"], paths)
        self.assertIn(["Staging", "Provisioning"], paths)

    def test_a_word_table_keeps_its_headings(self) -> None:
        with TemporaryDirectory() as tmp:
            text = extract(self._write_docx(Path(tmp)))[0].text
        self.assertIn("Stage: 2 | Coverage: 3%", text)

    def test_word_claims_no_page_numbers(self) -> None:
        # Word has no page concept until layout. A guessed page reference sends a reviewer
        # to the wrong part of the document and looks authoritative doing it.
        with TemporaryDirectory() as tmp:
            pages = extract(self._write_docx(Path(tmp)))
        self.assertEqual([p.page_number for p in pages], [1])


if __name__ == "__main__":
    unittest.main()
